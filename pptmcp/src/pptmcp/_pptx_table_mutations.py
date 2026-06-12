"""python-pptx backend for table mutation operations.

add_column       — append a new column to an existing table shape.
edit_table_cell  — write text to a table cell by row / column index.
set_column_width — set a column's width (in inches).

No FastMCP / MCP imports. No COM / pywin32 imports. Pure python-pptx + lxml + stdlib.
"""
from __future__ import annotations

import logging
from copy import deepcopy

from lxml import etree as ET
from pptx.oxml.ns import qn

from pptmcp.presentation_pptx import (
    _atomic_save,
    _audit_log,
    _check_confirm,
    _check_path,
    _check_write,
    _evict_prs,
    _load_prs,
)

logger = logging.getLogger(__name__)

_EMU_PER_INCH = 914400


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _find_table_shape(prs, slide_index: int, shape_id: int):  # noqa: ANN001
    """Return the table shape matching shape_id on slide_index; raise ValueError if not found."""
    try:
        slide = prs.slides[slide_index]
    except IndexError:
        raise ValueError(f"Slide index {slide_index} is out of range")
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            if not shape.has_table:
                raise ValueError(
                    f"Shape {shape_id} on slide {slide_index} is not a table"
                )
            return shape
    raise ValueError(f"No shape with id={shape_id} found on slide {slide_index}")


def _make_empty_tc() -> ET._Element:
    """Build a minimal empty <a:tc> element with a proper <a:txBody>."""
    tc = ET.Element(qn("a:tc"))
    tx_body = ET.SubElement(tc, qn("a:txBody"))
    ET.SubElement(tx_body, qn("a:bodyPr"))
    ET.SubElement(tx_body, qn("a:p"))
    return tc


def _clear_tc_text(tc: ET._Element) -> None:
    """Clear all text content from a <a:tc> element, leaving formatting intact."""
    for tx_body in tc.findall(qn("a:txBody")):
        for para in tx_body.findall(qn("a:p")):
            for run in para.findall(qn("a:r")):
                t_el = run.find(qn("a:t"))
                if t_el is not None:
                    t_el.text = ""


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def add_column(
    path: str,
    slide_index: int,
    shape_id: int,
    width_inches: float = 1.0,
    confirm: bool = False,
) -> dict:
    """Append a new column to a table shape, preserving row formatting.

    Gate 1 (env): PPT_ENABLE_WRITE=true.
    Gate 2 (param): confirm=True.

    Args:
        path: Presentation file path (must be in allowlist roots).
        slide_index: Zero-based slide index.
        shape_id: shape_id of the table shape.
        width_inches: Width of the new column in inches (default 1.0).
        confirm: Must be True to proceed.
    """
    _check_write()
    _check_confirm(confirm)

    resolved = _check_path(path)
    prs = _load_prs(resolved)
    shape = _find_table_shape(prs, slide_index, shape_id)
    table = shape.table
    tbl = table._tbl

    tbl_grid = tbl.find(qn("a:tblGrid"))
    if tbl_grid is None:
        raise ValueError("Table XML is malformed: missing tblGrid element")

    width_emu = max(1, int(width_inches * _EMU_PER_INCH))
    grid_col = ET.SubElement(tbl_grid, qn("a:gridCol"))
    grid_col.set("w", str(width_emu))

    for tr in tbl.findall(qn("a:tr")):
        tcs = tr.findall(qn("a:tc"))
        if tcs:
            new_tc = deepcopy(tcs[-1])
            _clear_tc_text(new_tc)
        else:
            new_tc = _make_empty_tc()
        tr.append(new_tc)

    col_count = len(tbl_grid.findall(qn("a:gridCol")))
    _audit_log("add_column", resolved, slide_index, {"shape_id": shape_id, "col_count": col_count})
    _atomic_save(prs, resolved)
    _evict_prs(resolved)
    return {
        "status": "column_added",
        "slide_index": slide_index,
        "shape_id": shape_id,
        "column_count": col_count,
        "width_inches": width_inches,
    }


def edit_table_cell(
    path: str,
    slide_index: int,
    shape_id: int,
    row: int,
    col: int,
    text: str,
    confirm: bool = False,
) -> dict:
    """Write text to a table cell by zero-based row and column index.

    Gate 1 (env): PPT_ENABLE_WRITE=true.
    Gate 2 (param): confirm=True.

    Args:
        path: Presentation file path.
        slide_index: Zero-based slide index.
        shape_id: shape_id of the table shape.
        row: Zero-based row index.
        col: Zero-based column index.
        text: Text to write into the cell.
        confirm: Must be True to proceed.
    """
    _check_write()
    _check_confirm(confirm)

    resolved = _check_path(path)
    prs = _load_prs(resolved)
    shape = _find_table_shape(prs, slide_index, shape_id)
    table = shape.table

    num_rows = len(table.rows)
    num_cols = len(table.columns)
    if row < 0 or row >= num_rows:
        raise ValueError(f"Row index {row} out of range (table has {num_rows} rows)")
    if col < 0 or col >= num_cols:
        raise ValueError(f"Column index {col} out of range (table has {num_cols} columns)")

    cell = table.cell(row, col)
    cell.text = text

    _audit_log("edit_table_cell", resolved, slide_index, {"shape_id": shape_id, "row": row, "col": col})
    _atomic_save(prs, resolved)
    _evict_prs(resolved)
    return {
        "status": "cell_updated",
        "slide_index": slide_index,
        "shape_id": shape_id,
        "row": row,
        "col": col,
        "text": text,
    }


def set_column_width(
    path: str,
    slide_index: int,
    shape_id: int,
    col: int,
    width_inches: float,
    confirm: bool = False,
) -> dict:
    """Set the width of a table column in inches.

    Gate 1 (env): PPT_ENABLE_WRITE=true.
    Gate 2 (param): confirm=True.

    Args:
        path: Presentation file path.
        slide_index: Zero-based slide index.
        shape_id: shape_id of the table shape.
        col: Zero-based column index.
        width_inches: New column width in inches.
        confirm: Must be True to proceed.
    """
    _check_write()
    _check_confirm(confirm)

    resolved = _check_path(path)
    prs = _load_prs(resolved)
    shape = _find_table_shape(prs, slide_index, shape_id)
    table = shape.table

    num_cols = len(table.columns)
    if col < 0 or col >= num_cols:
        raise ValueError(f"Column index {col} out of range (table has {num_cols} columns)")

    width_emu = max(1, int(width_inches * _EMU_PER_INCH))
    table.columns[col].width = width_emu

    _audit_log("set_column_width", resolved, slide_index, {"shape_id": shape_id, "col": col, "width_inches": width_inches})
    _atomic_save(prs, resolved)
    _evict_prs(resolved)
    return {
        "status": "column_width_set",
        "slide_index": slide_index,
        "shape_id": shape_id,
        "col": col,
        "width_inches": width_inches,
    }
